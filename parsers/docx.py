"""DOCX parser: python-docx → markdown.

Walks paragraphs (mapping Heading 1-6 styles to # ... ###### prefixes) and
tables (rendered as GitHub-flavored markdown). Numbered/bulleted list styles
become `-` / `1.` prefixes. Embedded images are not OCR'd in this pass —
DOCX archives are overwhelmingly text-bearing; if image-heavy docs land here
later we add a hybrid path.
"""

from __future__ import annotations

from typing import Iterator


def page_count(_path: str) -> int:
    return 1


def iter_pages(
    path: str,
    *,
    page_start: int = 1,
    page_end: int | None = None,
    staging_dir: str | None = None,
) -> Iterator[dict]:
    del page_start, page_end, staging_dir
    from docx import Document

    doc = Document(path)
    out: list[str] = []
    for block in _iter_blocks(doc):
        if block["kind"] == "p":
            md = _para_md(block["para"])
            if md:
                out.append(md)
        else:
            out.append(_table_md(block["table"]))
    yield {"page": 0, "markdown": "\n\n".join(out), "seconds": 0.0, "cached": False}


def _iter_blocks(doc):
    """Yield paragraphs + tables in document order."""
    from docx.oxml.ns import qn
    body = doc.element.body
    para_iter = iter(doc.paragraphs)
    table_iter = iter(doc.tables)
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield {"kind": "p", "para": next(para_iter, None)}
        elif child.tag == qn("w:tbl"):
            yield {"kind": "tbl", "table": next(table_iter, None)}


def _para_md(p) -> str:
    if p is None:
        return ""
    text = p.text.strip()
    if not text:
        return ""
    style = (p.style.name or "") if p.style else ""
    if style.startswith("Heading "):
        try:
            level = int(style.split(" ", 1)[1])
            level = max(1, min(6, level))
            return f"{'#' * level} {text}"
        except ValueError:
            pass
    if style in ("List Bullet", "List Bullet 2", "List Bullet 3"):
        return f"- {text}"
    if style in ("List Number", "List Number 2", "List Number 3"):
        return f"1. {text}"
    return text


def _table_md(tbl) -> str:
    if tbl is None or not tbl.rows:
        return ""
    rows = [[cell.text.replace("\n", " ").strip() for cell in row.cells] for row in tbl.rows]
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    header = rows[0]
    body = rows[1:] if len(rows) > 1 else []
    out = ["| " + " | ".join(header) + " |",
           "| " + " | ".join("---" for _ in header) + " |"]
    out.extend("| " + " | ".join(r) + " |" for r in body)
    return "\n".join(out)
